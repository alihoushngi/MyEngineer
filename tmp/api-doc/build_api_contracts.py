from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/ali/Desktop/Projects/MyEnginner")
OUTPUT = ROOT / "output" / "docx" / "mohandes-man-api-contracts-mvp-fa.docx"

FONT = "Tahoma"
MONO_FONT = "Courier New"
NAVY = "17365D"
PALE_BLUE = "F3F7FB"
LIGHT_BLUE = "DCE6F1"
MID_BLUE = "5B9BD5"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = "4A4A4A"
WHITE = "FFFFFF"
BLACK = "000000"


PHASE_1 = [
    ("F1-01", "مکان", "listCatalogCities", "بدون ورودی", "City[] شامل id و name و provinceId", "services/catalog-service/catalog-service.ts"),
    ("F1-02", "مکان", "getProvinces", "بدون ورودی", "Province[] شامل id و name", "services/city-service/city-service.ts"),
    ("F1-03", "مکان", "getCitiesByProvince", "provinceId", "City[] مربوط به استان", "services/city-service/city-service.ts"),
    ("F1-04", "جستجو", "searchCatalog", "query و cities[]", "SearchCatalogResult شامل query و services و experts", "services/search-service/search-service.ts"),
    ("F1-05", "خدمت", "getServiceDetail", "slug از شش ServiceSlug موجود", "ServiceDetailData یا null", "services/catalog-service/catalog-service.ts"),
    ("F1-06", "کشف متخصص", "قرارداد فهرست و فیلتر متخصصان خدمت", "serviceSlug و city و skill و experience و license و discipline و degree و tab و page", "ExpertCardData[]، تعداد نتیجه و داده صفحه بندی. شکل صفحه بندی در پروژه تعیین نشده است", "docs/PHASE-1-SCOPE.md و config/service-filters.config"),
    ("F1-07", "متخصص", "getExpertCardData", "id", "ExpertCardData یا null", "services/expert-service/expert-service.ts"),
    ("F1-08", "متخصص", "getExpertProfile", "id", "ExpertProfile یا null، شامل contact در صورت وجود", "services/expert-service/expert-service.ts"),
    ("F1-09", "کاربر", "getUserAccess", "نشست جاری", "UserAccessResult و UserSession با role برابر user", "services/user-auth-service/user-access-service.ts"),
    ("F1-10", "ورود کاربر", "requestUserLoginOtp", "phone", "موفقیت بدون بدنه یا ServiceMutationFailure", "services/user-auth-service/user-auth-service.ts"),
    ("F1-11", "ورود کاربر", "loginUserWithOtp", "phone و otp شش رقمی", "ایجاد نشست کاربر و موفقیت بدون بدنه", "services/user-auth-service/user-auth-service.ts"),
    ("F1-12", "ثبت نام کاربر", "requestUserRegisterOtp", "phone", "موفقیت بدون بدنه یا ServiceMutationFailure", "services/user-auth-service/user-auth-service.ts"),
    ("F1-13", "ثبت نام کاربر", "verifyUserRegisterOtp", "phone و otp شش رقمی", "موفقیت بدون بدنه", "services/user-auth-service/user-auth-service.ts"),
    ("F1-14", "ثبت نام کاربر", "completeUserRegister", "phone و otp و displayName حداقل دو نویسه و password حداقل هشت نویسه", "ایجاد نشست کاربر و موفقیت بدون بدنه", "services/user-auth-service/user-auth-service.ts"),
    ("F1-15", "نشست کاربر", "logoutUser", "نشست جاری", "لغو نشست و موفقیت بدون بدنه", "services/user-auth-service/user-auth-service.ts"),
    ("F1-16", "متخصص", "getEngineerAccess", "نشست جاری", "EngineerAccessResult و در حالت مجاز EngineerWorkspace", "services/engineer-service/engineer-access-service.ts"),
    ("F1-17", "ورود متخصص", "requestEngineerLoginOtp", "phone", "موفقیت بدون بدنه یا ServiceMutationFailure", "services/engineer-auth-service/engineer-auth-service.ts"),
    ("F1-18", "ورود متخصص", "loginEngineerWithOtp", "phone و otp شش رقمی", "ایجاد نشست متخصص و موفقیت بدون بدنه", "services/engineer-auth-service/engineer-auth-service.ts"),
    ("F1-19", "نشست متخصص", "logoutEngineer", "نشست جاری", "لغو نشست و موفقیت بدون بدنه", "services/engineer-auth-service/engineer-auth-service.ts"),
    ("F1-20", "ثبت نام متخصص", "sendOtp", "phone و nationalId", "موفقیت بدون بدنه. موبایل 11 رقم و کد ملی 10 رقم با checksum", "services/registration-service/registration-service.ts"),
    ("F1-21", "ثبت نام متخصص", "verifyOtp", "phone و code پنج رقمی", "موفقیت بدون بدنه", "services/registration-service/registration-service.ts"),
    ("F1-22", "ثبت نام متخصص", "getExpertiseCatalog", "بدون ورودی", "expertise[] و software[]، هر مورد شامل id و label", "services/registration-service/registration-service.ts"),
    ("F1-23", "ثبت نام متخصص", "saveServiceArea", "provinceId و cityId و nearbyCityIds[]", "موفقیت بدون بدنه", "services/registration-service/registration-service.types.ts"),
    ("F1-24", "ثبت نام متخصص", "saveExpertise", "expertiseIds[] و softwareIds[]", "موفقیت بدون بدنه", "services/registration-service/registration-service.types.ts"),
    ("F1-25", "ثبت نام متخصص", "savePersonalInfo", "firstName و lastName و avatarUploadId اختیاری", "موفقیت بدون بدنه", "services/registration-service/registration-service.types.ts"),
    ("F1-26", "آپلود", "قرارداد آپلود فایل ثبت نام", "فایل و نوع مصرف شامل avatar، degree، license، portfolioImage یا certificate", "uploadId رشته ای. محدودیت حجم در پروژه تعریف نشده است", "types/store/registration.types.ts و کامپوننت های FileUpload"),
    ("F1-27", "ثبت نام متخصص", "saveEducation", "level و degrees[] و degreeFileUploadIds", "موفقیت بدون بدنه", "services/registration-service/registration-service.types.ts"),
    ("F1-28", "ثبت نام متخصص", "saveOrganization", "isMember و membershipNumber و hasLicense و licenseNumber و licenseUploadId و discipline و qualifications[]", "موفقیت بدون بدنه", "services/registration-service/registration-service.types.ts"),
    ("F1-29", "ثبت نام متخصص", "saveResume", "experienceYears بزرگ تر یا مساوی صفر و resumeText حداقل ده نویسه", "موفقیت بدون بدنه", "services/registration-service/registration-service.types.ts"),
    ("F1-30", "ثبت نام متخصص", "submitRegistration", "imageCount و certificateCount و acceptRules برابر true و profile اختیاری در مرز mock فعلی", "ثبت نهایی و ایجاد نشست متخصص با source برابر registration", "services/registration-service/registration-service.ts"),
    ("F1-31", "درخواست خدمت", "createServiceRequest", "expertId و serviceSlug و cityId و description حداقل بیست نویسه", "requestId", "services/user-marketplace-service/user-marketplace-service.ts"),
    ("F1-32", "حساب کاربر", "getUserWorkspace", "نشست کاربر", "UserWorkspace یا null", "services/user-account-service/user-account-service.ts"),
    ("F1-33", "درخواست کاربر", "getUserRequest", "id و نشست کاربر", "UserRequest یا null. فهرست درخواست ها نیز در UserWorkspace مصرف می شود", "services/user-account-service/user-account-service.ts"),
    ("F1-34", "پنل متخصص", "getEngineerWorkspace", "نشست متخصص", "EngineerWorkspace یا null", "services/engineer-service/engineer-access-service.ts"),
    ("F1-35", "درخواست متخصص", "getEngineerRequest", "id و نشست متخصص", "EngineerRequest یا null. فهرست درخواست ها نیز در EngineerWorkspace مصرف می شود", "services/engineer-service/engineer-access-service.ts"),
]


PHASE_2 = [
    ("F2-01", "ورود کاربر", "loginUserWithPassword", "phone و password", "ایجاد نشست UserSession", "services/user-auth-service/user-auth-service.ts"),
    ("F2-02", "ورود متخصص", "loginEngineerWithPassword", "phone و password", "ایجاد نشست EngineerSession", "services/engineer-auth-service/engineer-auth-service.ts"),
    ("F2-03", "خانه", "getHomeCatalog", "بدون ورودی", "HomeCatalogData شامل experts، cities، popularServices، drawingServices، faqCategories و knowledgeTips", "services/catalog-service/catalog-service.ts"),
    ("F2-04", "ذخیره متخصص", "getCurrentSavedExpertIds", "نشست کاربر", "string[] از شناسه متخصصان", "services/user-account-service/user-account-service.ts"),
    ("F2-05", "ذخیره متخصص", "toggleSavedExpert", "expertId و نشست کاربر", "saved به صورت boolean", "services/user-marketplace-service/user-marketplace-service.ts"),
    ("F2-06", "پیام", "listConversations", "نقش و نشست جاری، صفحه فهرست", "فهرست Conversation برای همان کاربر یا متخصص. قرارداد مستقل در مستند پروژه خواسته شده است", "docs/MESSAGING.md"),
    ("F2-07", "پیام", "getConversation و پیام های آن", "conversationId و cursor اختیاری برای پیام های قدیمی تر", "Conversation و Message[]. تاریخچه با cursor است و شماره صفحه برای پیام ها تعریف نشده است", "docs/MESSAGING.md و getUserMessages و getEngineerMessages"),
    ("F2-08", "پیام", "startOrOpenConversation", "expertId و نشست کاربر", "conversationId. قاعده یکتایی گفتگو هنوز تصمیم محصول می خواهد", "services/messaging-service/messaging-service.ts"),
    ("F2-09", "پیام", "sendMessage و sendEngineerMessage", "conversationId و body غیرخالی", "موفقیت بدون بدنه و پیام ذخیره شده با status برابر sent", "services/messaging-service/messaging-service.ts"),
    ("F2-10", "پیام", "markConversationRead", "conversationId و نشست جاری", "موفقیت بدون بدنه", "services/messaging-service/messaging-service.ts"),
    ("F2-11", "نظر خدمت", "فهرست و جزئیات نظر برای دو نقش", "نشست و reviewId اختیاری", "UserReviewItem برای کاربر و EngineerReview برای متخصص", "getUserReview و getEngineerReview"),
    ("F2-12", "نظر خدمت", "قرارداد review eligibility", "کاربر جاری و requestId و رابطه درخواست با متخصص", "مجاز یا غیرمجاز بودن ثبت نظر. شکل پاسخ در پروژه تعیین نشده است", "docs/USER-ACCOUNT.md و docs/ENGINEER-PANEL.md"),
    ("F2-13", "نظر خدمت", "submitReview", "requestId و rating عدد صحیح 1 تا 5 و body بین 10 تا 2000 نویسه", "reviewId", "services/review-service/review-service.ts"),
    ("F2-14", "اعلان", "فهرست اعلان های نقش جاری", "نشست کاربر یا متخصص", "AppNotification[] فقط برای همان recipientRole و recipientId", "UserWorkspace و getEngineerNotifications"),
    ("F2-15", "اعلان", "markNotificationRead", "notificationId و نشست جاری", "موفقیت بدون بدنه", "services/notification-service/notification-service.ts"),
    ("F2-16", "پروفایل متخصص", "updateEngineerProfile", "firstName و lastName و profession و about اختیاری", "موفقیت بدون بدنه", "services/engineer-service/engineer-service.ts"),
    ("F2-17", "پروفایل متخصص", "updateEngineerSpecialties", "specialties[] و software[]", "موفقیت بدون بدنه", "services/engineer-service/engineer-service.ts"),
    ("F2-18", "مکان متخصص", "getEngineerLocationCatalog", "نشست متخصص", "provinces[] و cities[]", "services/engineer-service/engineer-service.ts"),
    ("F2-19", "مکان متخصص", "updateEngineerServiceArea", "provinceId و cityId و nearbyCityIds[]", "موفقیت بدون بدنه", "services/engineer-service/engineer-service.ts"),
    ("F2-20", "نمونه کار", "addEngineerPortfolioItem", "title و description اختیاری. UI فایل تصویر هم می گیرد ولی service فعلی آن را ارسال نمی کند", "موفقیت بدون بدنه", "services/engineer-service/engineer-service.ts"),
    ("F2-21", "نمونه کار", "removeEngineerPortfolioItem", "id", "موفقیت بدون بدنه", "services/engineer-service/engineer-service.ts"),
    ("F2-22", "نمونه کار", "قرارداد خواندن و ویرایش و آپلود نمونه کار", "شناسه آیتم، metadata و uploadId", "ExpertPortfolioItem یا فهرست آن. ویرایش و آپلود در مستند پنل صریح است ولی service نهایی نشده است", "docs/ENGINEER-PANEL.md"),
    ("F2-23", "مدارک", "قرارداد خواندن و ویرایش و آپلود مدارک", "kind و metadata و uploadId", "EngineerCredential[] با وضعیت submitted، pending_review، verified یا needs_correction", "docs/ENGINEER-PANEL.md و types/store/engineer.types.ts"),
]


PHASE_3 = [
    ("F3-01", "مقاله", "listArticles", "بدون ورودی", "ArticleCardData[]", "services/article-service/article-service.ts"),
    ("F3-02", "مقاله", "listArticleCategories", "بدون ورودی", "ArticleCategory[]", "services/article-service/article-service.ts"),
    ("F3-03", "مقاله", "getArticleCategory", "slug", "ArticleCategory یا null", "services/article-service/article-service.ts"),
    ("F3-04", "مقاله", "listArticlesByCategory", "slug", "ArticleCardData[]", "services/article-service/article-service.ts"),
    ("F3-05", "مقاله", "getArticleBySlug", "slug", "Article یا null", "services/article-service/article-service.ts"),
    ("F3-06", "نظر مقاله", "listArticleComments", "articleId", "ArticleComment[]", "services/article-comment-service/article-comment-service.ts"),
    ("F3-07", "نظر مقاله", "submitArticleComment", "articleId و authorName حداقل دو نویسه و phone معتبر و body حداقل ده نویسه", "ArticleComment", "services/article-comment-service/article-comment-service.ts"),
    ("F3-08", "پرسش متداول", "listFaqCategories", "بدون ورودی", "FaqCategory[]", "services/faq-service/faq-service.ts"),
    ("F3-09", "پرسش متداول", "getFaqCategory", "slug", "FaqCategoryDetail یا null", "services/faq-service/faq-service.ts"),
    ("F3-10", "دانش", "listKnowledgeCategories", "بدون ورودی", "KnowledgeCategory[]", "services/knowledge-service/knowledge-service.ts"),
    ("F3-11", "دانش", "getKnowledgeCategory", "slug", "KnowledgeCategoryDetail یا null", "services/knowledge-service/knowledge-service.ts"),
]


MODELS = [
    ("Province", "id: string، name: string", "types/store/registration.types.ts"),
    ("City", "id: string، name: string، provinceId: string", "types/store/registration.types.ts"),
    ("ExpertCardData", "id، href، name، profession؛ فیلدهای اختیاری avatarSrc، primarySpecialty، city، experienceYears، isVerified، isActive، rating، reviewCount، specialties، serviceSlugs، discipline، degree، hasLicense، track", "types/store/expert.types.ts"),
    ("ExpertProfile", "id، name، profession؛ فیلدهای اختیاری معرفی، درباره، وضعیت، بازدید، سابقه، شهرها، تخصص ها، نرم افزار، تحصیلات، عضویت، پروانه، صلاحیت، گواهی، تاریخچه، نمونه کار، امتیاز، نظر، contact و relatedExperts", "types/store/expert.types.ts"),
    ("ServiceDetailData", "slug، title، eyebrow، description، longDescription، imageSrc، imageAlt، accent، specialties، process، faqs، experts؛ scopeItems و showSuggestedExperts اختیاری", "types/store/service.types.ts"),
    ("SearchCatalogResult", "query: string، services: ServiceCategory[]، experts: ExpertCardData[]", "types/store/search.types.ts"),
    ("HomeCatalogData", "experts، cities، popularServices، drawingServices، faqCategories، knowledgeTips", "types/store/home.types.ts"),
    ("UserSession", "isAuthenticated: true، role: user، isMock، source: login یا registration، profile اختیاری", "types/store/user-auth.types.ts"),
    ("EngineerSession", "isAuthenticated: true، role: engineer، isMock، source: login یا registration، profile اختیاری", "types/store/engineer-auth.types.ts"),
    ("UserAccessResult", "checking، authenticated، engineer_session، unauthenticated، expired، unavailable یا error", "types/store/user-auth.types.ts"),
    ("EngineerAccessResult", "visual_review، unavailable، unauthenticated، forbidden، registration_in_progress، pending_review یا active؛ حالت مجاز شامل workspace است", "types/store/engineer.types.ts"),
    ("CreateServiceRequestInput", "expertId، serviceSlug، cityId، description", "types/store/service-request.types.ts"),
    ("ServiceRequest", "id، title، serviceSlug، serviceLabel، city، cityId، createdAtLabel، summary، description، status، expertId، expertName، expertHref، customerId، customerDisplayName، conversationId", "types/store/service-request.types.ts"),
    ("UserWorkspace", "account، requests، conversations، messagesByConversationId، savedExperts، reviews، notifications", "types/store/user-account.types.ts"),
    ("EngineerWorkspace", "account، profile، services، serviceArea، requests، conversations، messagesByConversationId، portfolio، credentials، reviews، notifications", "types/store/engineer.types.ts"),
    ("Conversation", "id، participants، relatedRequestId، relatedServiceLabel، relatedEngineerId، relatedCustomerId، latestMessage، unreadByRole، createdAtLabel، updatedAtLabel، updatedAtMs", "types/store/messaging.types.ts"),
    ("Message", "id، conversationId، senderRole، senderId، content، createdAtLabel، createdAtMs، status: sent", "types/store/messaging.types.ts"),
    ("AppNotification", "id، recipientRole، recipientId، kind، title، body، createdAtLabel، createdAtMs، isRead، href", "types/store/notification.types.ts"),
    ("ServiceReview", "id، expertId، expertName، authorCustomerId، authorDisplayName، relatedRequestId، relatedServiceLabel، rating، text، dateLabel، createdAtMs؛ highlights و replyText اختیاری", "types/store/review.types.ts"),
    ("Article", "ArticleCardData به همراه body، viewCount، faqs، relatedServiceHref و relatedServiceLabel اختیاری", "types/store/article.types.ts"),
    ("FaqCategoryDetail", "FaqCategory به همراه items و relatedCategories اختیاری", "types/store/faq.types.ts"),
    ("KnowledgeCategoryDetail", "KnowledgeCategory به همراه tips", "types/store/knowledge.types.ts"),
    ("ServiceMutationFailure", "ok: false، message، status و code یکی از unavailable، unauthorized، validation یا server", "types/store/engineer-auth.types.ts"),
]


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "1")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), LIGHT_GRAY)


def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    bidi = tbl_pr.find(qn("w:bidiVisual"))
    if bidi is None:
        bidi = OxmlElement("w:bidiVisual")
        tbl_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_paragraph_rtl(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run_style(run, size=10.5, bold=False, color=BLACK, rtl=True, mono=False):
    font_name = MONO_FONT if mono else FONT
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:cs"), FONT)
    r_fonts.set(qn("w:eastAsia"), font_name)
    rtl_node = r_pr.find(qn("w:rtl"))
    if rtl and rtl_node is None:
        rtl_node = OxmlElement("w:rtl")
        rtl_node.set(qn("w:val"), "1")
        r_pr.append(rtl_node)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:bidi"), "fa-IR")
    lang.set(qn("w:val"), "fa-IR")


def add_para(doc, text="", style=None, bold_lead=None, before=0, after=6, keep=False):
    p = doc.add_paragraph(style=style)
    set_paragraph_rtl(p)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_style(lead, bold=True)
        body = p.add_run(text[len(bold_lead):])
        set_run_style(body)
    else:
        run = p.add_run(text)
        set_run_style(run)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    set_run_style(p.add_run(text), size=10)
    return p


def clear_cell(cell):
    p = cell.paragraphs[0]
    for run in list(p.runs):
        p._element.remove(run._element)
    return p


def write_cell(cell, text, size=8.2, bold=False, color=BLACK, code=False, center=False):
    p = clear_cell(cell)
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    set_run_style(p.add_run(str(text)), size=size, bold=bold, color=color, rtl=not code, mono=code)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_api_field(doc, label, value, *, keep=True, muted=False):
    if muted:
        value = value.replace("/", "/\u200b").replace("-", "-\u200b")
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_rtl(p)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    set_run_style(p.add_run(f"{label}: "), size=9.6, bold=True, color=TEXT_GRAY if muted else BLACK)
    set_run_style(p.add_run(value), size=9.6 if not muted else 8.8, color=TEXT_GRAY if muted else BLACK)
    return p


def normalize_request(inputs):
    if inputs == "بدون ورودی":
        return "نیاز ندارد"
    if inputs == "نشست جاری":
        return "Body نیاز ندارد؛ نشست جاری باید همراه درخواست ارسال شود"
    if inputs == "نشست کاربر":
        return "Body نیاز ندارد؛ نشست معتبر کاربر باید همراه درخواست ارسال شود"
    if inputs == "نشست متخصص":
        return "Body نیاز ندارد؛ نشست معتبر متخصص باید همراه درخواست ارسال شود"
    return inputs


def add_api_list(doc, rows):
    for api_id, domain, contract, inputs, output, _source in rows:
        heading = add_heading(doc, f"{api_id}  {domain}", 3)
        heading.paragraph_format.space_before = Pt(9)
        heading.paragraph_format.space_after = Pt(2)
        heading.paragraph_format.keep_with_next = True

        contract_p = doc.add_paragraph()
        set_paragraph_rtl(contract_p)
        contract_p.paragraph_format.space_after = Pt(3)
        contract_p.paragraph_format.keep_with_next = True
        set_run_style(contract_p.add_run("قرارداد فرانت: "), size=9.6, bold=True)
        set_run_style(contract_p.add_run(contract), size=9.2, mono=contract.isascii(), rtl=not contract.isascii())

        add_api_field(doc, "Request به بک اند", normalize_request(inputs))
        response_p = add_api_field(doc, "Response از بک اند", output, keep=False)
        response_p.paragraph_format.space_after = Pt(7)


def add_model_list(doc, rows):
    for model, fields, _source in rows:
        heading = add_heading(doc, model, 3)
        heading.paragraph_format.space_before = Pt(8)
        heading.paragraph_format.space_after = Pt(2)
        fields_p = add_api_field(doc, "فیلدهای Response", fields, keep=False)
        fields_p.paragraph_format.space_after = Pt(6)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)

    for style_name, size, spacing in (("Title", 23, 10), ("Subtitle", 12, 16), ("Heading 1", 16, 9), ("Heading 2", 12.5, 6), ("Heading 3", 11, 4)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(BLACK if style_name != "Subtitle" else TEXT_GRAY)
        style.paragraph_format.space_before = Pt(spacing)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:cs"), FONT)
        if style_name == "Title":
            p_pr = style._element.get_or_add_pPr()
            p_borders = p_pr.find(qn("w:pBdr"))
            if p_borders is not None:
                p_pr.remove(p_borders)


def add_footer(section):
    p = section.footer.paragraphs[0]
    set_paragraph_rtl(p, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(4)
    set_run_style(p.add_run("مهندس من  |  فهرست قراردادهای API  |  صفحه "), size=8.3, color=TEXT_GRAY)
    run = p.add_run()
    set_run_style(run, size=8.3, color=TEXT_GRAY)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr)
    run._r.append(fld_char_2)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_paragraph_rtl(p)
    for run in p.runs:
        set_run_style(run, size={1: 16, 2: 12.5, 3: 11}.get(level, 10.5), bold=True)
    return p


def add_summary_list(doc):
    add_bullet(doc, f"فاز 1 شامل {len(PHASE_1)} قرارداد ضروری است و شرط انتشار MVP محسوب می شود")
    add_bullet(doc, f"فاز 2 شامل {len(PHASE_2)} قرارداد مهم است و بعد از پایدار شدن فاز 1 انجام می شود")
    add_bullet(doc, f"فاز 3 شامل {len(PHASE_3)} قرارداد محتوا و پشتیبانی است و انتشار دوهفته ای را متوقف نمی کند")


def build_document():
    doc = Document()
    style_document(doc)
    add_footer(doc.sections[0])
    doc.core_properties.title = "فهرست قراردادهای API پروژه مهندس من"
    doc.core_properties.subject = "اولویت بندی سه فازی برای MVP دوهفته ای"
    doc.core_properties.author = "تیم پروژه مهندس من"

    title = doc.add_paragraph(style="Title")
    set_paragraph_rtl(title, WD_ALIGN_PARAGRAPH.CENTER)
    title_p_pr = title._p.get_or_add_pPr()
    title_borders = title_p_pr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_p_pr.remove(title_borders)
    set_run_style(title.add_run("فهرست قراردادهای API پروژه مهندس من"), size=23, bold=True)

    subtitle = doc.add_paragraph(style="Subtitle")
    set_paragraph_rtl(subtitle, WD_ALIGN_PARAGRAPH.CENTER)
    set_run_style(subtitle.add_run("اولویت بندی سه فازی برای انتشار MVP در دو هفته"), size=12, color=TEXT_GRAY)

    meta = doc.add_paragraph()
    set_paragraph_rtl(meta, WD_ALIGN_PARAGRAPH.CENTER)
    meta.paragraph_format.space_after = Pt(18)
    set_run_style(meta.add_run("تاریخ بررسی 2026-09-13"), size=9.5, color=TEXT_GRAY)

    add_para(
        doc,
        "نتیجه اصلی: برای هدف دوهفته ای، فقط فاز 1 باید شرط انتشار باشد. فاز 2 تعاملات مهم موجود در پروژه را کامل می کند و فاز 3 برای محتوای خواندنی و نظر مقاله است. این فهرست 69 قرارداد منطقی را از سرویس ها، مدل های TypeScript و مستندات همین مخزن استخراج کرده است.",
        after=10,
    )
    add_summary_list(doc)

    add_heading(doc, "قاعده دامنه", 2)
    add_bullet(doc, "هیچ مسیر URL یا نام endpoint در این سند ساخته نشده است، چون پروژه صریحا تعریف مسیرها را به قرارداد بک اند واگذار کرده است.")
    add_bullet(doc, "در هر مورد، Request فقط داده موردنیاز را مشخص می کند. جای آن در path یا query یا body و همچنین متد HTTP در پروژه تعیین نشده است.")
    add_bullet(doc, "نام قراردادها از توابع سرویس فرانت گرفته شده است. مواردی که هنوز تابع نهایی ندارند فقط وقتی آمده اند که مستندات یا کامنت کد عبارت API CONTRACT REQUIRED داشته اند.")
    add_bullet(doc, "بک اند می تواند چند قرارداد منطقی را در یک endpoint ترکیب کند، به شرط آنکه همه ورودی ها، خروجی ها، نقش ها و حالت های خطای ثبت شده پوشش داده شوند.")

    doc.add_page_break()
    add_heading(doc, "مبنای بررسی", 1)
    add_para(doc, "منبع اصلی، پوشه های services و types/store و فایل های validation و مستندات محصول داخل docs است. پروژه در حال حاضر بیشتر سرویس ها را با mock یا داده محلی اجرا می کند و NEXT_PUBLIC_API_BASE_URL هنوز به مسیرهای واقعی متصل نشده است.")
    add_para(doc, "شش خدمت موجود در دامنه عبارت اند از land-surveying، construction-workers، drawing، interior-design، building-permit و administrative-services. مورد تازه ای به taxonomy اضافه نشده است.")

    add_heading(doc, "قرارداد سراسری پاسخ و خطا", 2)
    add_bullet(doc, "مرز HTTP موجود JSON است و متدهای GET، POST، PUT، PATCH و DELETE را پشتیبانی می کند، اما انتخاب متد برای هر قرارداد در پروژه تعیین نشده است.")
    add_bullet(doc, "timeout پیش فرض فرانت 15000 میلی ثانیه است. پاسخ 204 یا بدنه خالی به void تبدیل می شود.")
    add_bullet(doc, "کدهای خطای شناخته شده عبارت اند از unconfigured، unavailable، network، timeout، aborted، validation، unauthorized، forbidden، not_found، conflict، server و unknown.")
    add_bullet(doc, "وضعیت های 400 و 422 به validation، وضعیت 401 به unauthorized، وضعیت 403 به forbidden، وضعیت 404 به not_found، وضعیت های 408 و 504 به timeout، وضعیت 409 به conflict و وضعیت های 500 به بالا به server نگاشت می شوند.")
    add_bullet(doc, "نشست کاربر و متخصص باید نقش جدا داشته باشد. برخورد نقش اشتباه با مسیر خصوصی باید 401 و 403 و 409 را طبق وضعیت واقعی متمایز کند.")

    doc.add_page_break()
    add_heading(doc, "فاز 1 قراردادهای ضروری برای انتشار MVP", 1)
    add_para(doc, "این فاز مسیر اصلی محصول را باز می کند: کشف متخصص، مشاهده پروفایل و تماس، ثبت نام کاربر و متخصص، ارسال درخواست خدمت و دیدن درخواست در دو پنل. اگر زمان کم شد، حذف قراردادهای این فاز باعث ناقص شدن جریان اصلی می شود.")
    add_api_list(doc, PHASE_1)

    add_heading(doc, "فاز 2 قراردادهای مهم برای تکمیل تعامل", 1)
    add_para(doc, "این فاز قابلیت های موجود در رابط فعلی را از حالت mock خارج می کند: ورود با رمز، صفحه خانه پویا، ذخیره متخصص، پیام، نظر خدمت، اعلان و ویرایش پنل متخصص. مستند Phase 1 خود پروژه، پیام و ذخیره و ارسال نظر را نیازمند تایید محصول دانسته است؛ بنابراین این فاز نباید انتشار اولیه را متوقف کند.")
    add_api_list(doc, PHASE_2)

    add_heading(doc, "فاز 3 قراردادهای محتوا و پشتیبانی", 1)
    add_para(doc, "این فاز مسیرهای خواندنی مقاله، پرسش متداول و دانش را به منبع واقعی وصل می کند. نسخه MVP می تواند تا آماده شدن این قراردادها از محتوای محلی موجود یا حالت خالی صادقانه استفاده کند.")
    add_api_list(doc, PHASE_3)

    add_heading(doc, "مدل های داده موردانتظار", 1)
    add_para(doc, "نام و فیلدهای زیر عینا از مدل های TypeScript پروژه آمده اند. فیلدهای اختیاری باید واقعا در داده موجود باشند؛ فرانت برای پر کردن جای خالی، مقدار ساختگی تولید نمی کند.")
    add_model_list(doc, MODELS)

    add_heading(doc, "قواعد امنیت و رفتار", 1)
    add_bullet(doc, "APIهای account و engineer و پیام و اعلان باید خصوصی و بدون cache عمومی باشند. service worker پروژه مسیرهای خصوصی و API را cache نمی کند.")
    add_bullet(doc, "فهرست گفتگو، جزئیات گفتگو، پیام ها و اعلان ها باید با actor فعلی محدود شوند. کاربر نباید به گفتگو یا اعلان کاربر دیگر دسترسی داشته باشد.")
    add_bullet(doc, "هویت مشتری در پیام فقط displayName است و شماره موبایل کامل یا ماسک شده نباید به عنوان نام چت برگردد.")
    add_bullet(doc, "رابط پیام فعلی request و response است. realtime، typing، presence، push و attachment در قرارداد فعلی وجود ندارد.")
    add_bullet(doc, "وضعیت پیام در مدل فعلی فقط sent است. رسید خواندن برای هر پیام تعریف نشده و mark read در سطح گفتگو است.")
    add_bullet(doc, "اعلان ها باید با recipientRole و recipientId جدا شوند و در دو پنل مخلوط نشوند.")
    add_bullet(doc, "شماره تماس و SMS پروفایل متخصص فقط وقتی نمایش داده می شود که API مقدار واقعی بدهد.")

    add_heading(doc, "تصمیم های باز که نباید به API تبدیل شوند", 1)
    add_para(doc, "موارد زیر در خود پروژه صریحا تصمیم محصول یا قرارداد نامشخص محسوب شده اند. تا زمان تصمیم، ساخت endpoint برای آن ها خارج از این فهرست است.")
    for item in [
        "مسیر URL و متد دقیق هر قرارداد",
        "قبول یا رد درخواست، اعلام قیمت و state machine کامل درخواست",
        "یک گفتگو برای هر جفت کاربر و متخصص یا یک گفتگو برای هر درخواست",
        "draft ثبت نام، امکان پرش بین مراحل و مقصد بعد از ثبت نهایی متخصص",
        "شعاع و حداکثر تعداد شهرهای نزدیک",
        "فعال یا غیرفعال کردن یک خدمت توسط متخصص",
        "reply متخصص به نظر، تنظیمات اعلان و مرتب سازی نمونه کار",
        "دانلود فرم های مهندسی، چون صفحه فعلی stub است و فایل یا metadata ندارد",
        "اشتراک گذاری پروفایل، چون می تواند کاملا client-side باشد و API آن تایید نشده است",
        "password reset و social login، چون در سطح محصول فعلی وجود ندارند",
        "realtime، typing، presence و attachment در پیام",
        "پرداخت، کارمزد، شکایت، حذف حساب و cookie consent که فقط در متن حقوقی یا backlog آمده اند",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "فهرست منابع داخل پروژه", 1)
    for item in [
        "docs/PHASE-1-SCOPE.md",
        "docs/PRODUCT-FLOWS.md",
        "docs/USER-AUTH.md",
        "docs/USER-ACCOUNT.md",
        "docs/ENGINEER-PANEL.md",
        "docs/MESSAGING.md",
        "services/** و types/store/**",
        "lib/validation/** و config/services.config و config/service-filters.config",
    ]:
        add_bullet(doc, item)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_document())
